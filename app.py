
import csv
import html
import io
from functools import lru_cache

from flask import Flask, render_template, request, jsonify, send_file

import config

app = Flask(__name__)


@lru_cache(maxsize=1)
def get_tagger():
    from predictor import POSTagger

    return POSTagger(config.MODEL_PATH, config.VOCAB_PATH)


def _max_token_count(tagged_sents):
    return max((len(s) for s in tagged_sents if s), default=0)


def _export_grid_header_rows(tagged_sents):
    """
    每句占两行：第一行 sentence + 各词各占一列；第二行同一句号 + 各列对应该词词性。
    列数 = 1 + 本批最长句词数（短句右侧空单元格补齐）。
    """
    mx = _max_token_count(tagged_sents)
    header = ["sentence"] + [f"w{i}" for i in range(1, mx + 1)]
    rows = []
    for si, sent in enumerate(tagged_sents, start=1):
        if not sent:
            continue
        pad = mx - len(sent)
        words_row = [si] + [w for w, _ in sent] + [""] * pad
        tags_row = [si] + [t for _, t in sent] + [""] * pad
        rows.append(words_row)
        rows.append(tags_row)
    return header, rows


def _interlinear_sentence_blocks(tagged_sents, between_cols="    ", extra_per_col=6):
    """生成每句两行（词行、标签行）对齐块，供 HTML 等宽视图使用。"""
    chunks = []
    for sn, sent in enumerate(tagged_sents, start=1):
        if not sent:
            continue
        col_w = [max(len(w), len(t), 1) + extra_per_col for w, t in sent]
        line_words = between_cols.join(w.ljust(col_w[i]) for i, (w, _) in enumerate(sent))
        line_tags = between_cols.join(t.ljust(col_w[i]) for i, (_, t) in enumerate(sent))
        chunks.append((sn, line_words.rstrip(), line_tags.rstrip()))
    return chunks


def _format_html_interlinear(tagged_sents) -> str:
    """单页 HTML：<pre> 强制等宽，在任意浏览器中上下列对齐，不依赖记事本字体。"""
    between_cols = "    "
    extra_per_col = 6
    bodies = []
    for sn, lw, lt in _interlinear_sentence_blocks(tagged_sents, between_cols, extra_per_col):
        pair = html.escape(lw + "\n" + lt, quote=False)
        bodies.append(f"<h2 class=\"sn\">【句 {sn}】</h2>\n<pre class=\"align\">{pair}</pre>\n")
    body_inner = "\n".join(bodies)
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>词性标注 — 等宽对齐视图</title>
<style>
  body {{ font-family: "Segoe UI", "Microsoft YaHei", sans-serif; margin: 1.25rem; color: #222; background: #fafafa; }}
  h1 {{ font-size: 1.15rem; margin: 0 0 0.75rem; }}
  p.note {{ font-size: 0.9rem; color: #555; margin: 0 0 1rem; line-height: 1.5; }}
  h2.sn {{ font-size: 0.95rem; margin: 1.1rem 0 0.35rem; color: #1565c0; }}
  pre.align {{
    font-family: Consolas, "Courier New", "NSimSun", "新宋体", "SimSun", monospace;
    font-size: 14px;
    line-height: 1.45;
    margin: 0 0 0.5rem;
    padding: 10px 14px;
    background: #fff;
    border: 1px solid #ccc;
    border-radius: 6px;
    overflow-x: auto;
    white-space: pre;
    tab-size: 4;
  }}
</style>
</head>
<body>
<h1>英文词性标注</h1>
<p class="note">本页已用 CSS 指定等宽字体，直接上下对照即可。</p>
{body_inner}
</body>
</html>
"""


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/health")
def health():
    import os

    ok = os.path.isfile(config.MODEL_PATH) and os.path.isfile(config.VOCAB_PATH)
    return jsonify(model_ready=ok)


@app.route("/api/tag", methods=["POST"])
def api_tag():
    data = request.get_json(force=True, silent=True) or {}
    text = (data.get("text") or "").strip()
    if not text:
        return jsonify(error="empty text", sentences=[]), 400
    try:
        tagger = get_tagger()
    except Exception as e:
        return jsonify(error=f"model load failed: {e}", sentences=[]), 503

    from predictor import tag_text_batch

    try:
        sents = tag_text_batch(tagger, text)
    except Exception as e:
        return jsonify(error=f"tagging failed: {e}", sentences=[]), 500
    out = [[{"word": w, "tag": t} for w, t in s] for s in sents]
    return jsonify(sentences=out)


@app.route("/api/upload", methods=["POST"])
def api_upload():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(error="no file"), 400
    name = f.filename.lower()
    raw = f.read()
    try:
        if name.endswith(".txt"):
            text = raw.decode("utf-8", errors="replace")
        elif name.endswith(".docx"):
            from docx import Document

            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        else:
            return jsonify(error="仅支持 .txt 或 .docx"), 400
    except Exception as e:
        return jsonify(error=str(e)), 400
    source = "TXT" if name.endswith(".txt") else "DOCX"
    return jsonify(text=text, filename=f.filename, source=source)


@app.route("/api/export", methods=["POST"])
def api_export():
    data = request.get_json(force=True, silent=True) or {}
    text = (data.get("text") or "").strip()
    fmt = (data.get("format") or "html").strip().lower()
    if fmt not in ("csv", "xlsx", "html"):
        return jsonify(error="format 须为 csv、xlsx 或 html"), 400
    if not text:
        return jsonify(error="empty text"), 400
    try:
        tagger = get_tagger()
    except Exception as e:
        return jsonify(error=f"model load failed: {e}"), 503

    from predictor import tag_text_batch

    try:
        sents = tag_text_batch(tagger, text)
    except Exception as e:
        return jsonify(error=f"tagging failed: {e}"), 500
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        header, rows = _export_grid_header_rows(sents)
        writer.writerow(header)
        for row in rows:
            writer.writerow(row)
        csv_bytes = "\ufeff" + buf.getvalue()
        mem = io.BytesIO(csv_bytes.encode("utf-8"))
        mem.seek(0)
        return send_file(
            mem,
            as_attachment=True,
            download_name="pos_tags.csv",
            mimetype="text/csv; charset=utf-8",
        )
    elif fmt == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except ImportError:
            return jsonify(error="openpyxl not installed"), 500
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "词性标注"
        header, rows = _export_grid_header_rows(sents)
        ws.append(header)
        for row in rows:
            ws.append(row)

        header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
        header_font = Font(bold=True, size=10)
        thin = Side(style="thin", color="B4B4B4")
        cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
        zebra_a = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        zebra_b = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")

        ncol = ws.max_column or 1
        max_row = ws.max_row or 1

        for ci in range(1, ncol + 1):
            c = ws.cell(row=1, column=ci)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = cell_border
            lett = get_column_letter(ci)
            if ci == 1:
                ws.column_dimensions[lett].width = 10
            else:
                ws.column_dimensions[lett].width = 14

        for ri in range(2, max_row + 1):
            sid = ws.cell(row=ri, column=1).value
            fill = (
                zebra_a
                if (isinstance(sid, int) and sid % 2 == 1)
                else zebra_b
                if isinstance(sid, int)
                else zebra_a
            )
            for ci in range(1, ncol + 1):
                cell = ws.cell(row=ri, column=ci)
                cell.border = cell_border
                cell.fill = fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        ws.freeze_panes = "A2"
        ws.print_title_rows = "1:1"

        mem = io.BytesIO()
        wb.save(mem)
        mem.seek(0)
        return send_file(
            mem,
            as_attachment=True,
            download_name="pos_tags.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        mem = io.BytesIO(_format_html_interlinear(sents).encode("utf-8"))
        mem.seek(0)
        return send_file(
            mem,
            as_attachment=True,
            download_name="pos_tags_aligned.html",
            mimetype="text/html; charset=utf-8",
        )


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)
